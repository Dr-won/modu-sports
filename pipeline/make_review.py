"""전문가 검수표(엑셀)를 만든다. 자동 분류 결과를 미리 채우고, 애매한 줄은 '먼저 봐 주세요'로 표시한다.

검수가 끝난 엑셀은 read_review.py 로 pipeline/review/classification.csv 로 바꾸면 매주 자동 갱신에 반영된다.
사용: python make_review.py <RAW_DIR> <엑셀 저장 경로>
"""
import collections, json, os, re, sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from classify import EXERCISE, auto_social, auto_welfare  # noqa: E402

RAW, OUT = sys.argv[1], sys.argv[2]
AGE = {'': '제한 없음', 'child': '아동·청소년', 'adult': '성인'}
DEG = {'': '무관', 'severe': '심한 장애만', 'mild': '심하지 않은 장애만'}
TG = lambda t: '전체' if t == '*' else t.replace('|', ', ')
HEAD = ['먼저 봐 주세요', '구분', '키', '서비스명', '제공기관·지역', '자동_포함', '자동_대상장애유형', '자동_나이', '자동_운동재활', '자동_장애정도',
        '확정_포함', '확정_대상장애유형', '확정_나이', '확정_운동재활', '확정_장애정도', '근거·메모', '요약', '지원대상 원문']

wb = Workbook()
guide = wb.active
guide.title = '사용법'
for line in [
    '모두의 특수체육 — 장애 관련 서비스 분류 검수표',
    '',
    '1. "사회서비스"와 "복지로" 두 시트가 있습니다. 자동 분류 결과가 "자동_" 칸에 채워져 있습니다.',
    '2. 자동 분류가 맞으면 아무것도 적지 않아도 됩니다. 틀린 줄만 "확정_" 칸에서 골라 주세요(칸을 누르면 목록이 나옵니다).',
    '3. "먼저 봐 주세요"에 표시된 줄이 애매한 것들입니다. 시간이 없으면 이 줄만 보셔도 됩니다.',
    '4. 대상장애유형은 여러 개면 쉼표로: 예) 뇌병변, 지체   /  모든 장애유형이면 "전체"',
    '5. 다 보시면 파일을 저장해서 알려 주세요. 그 뒤로는 매주 자동 갱신 때 이 표를 따릅니다.',
    '',
    '칸 설명',
    '- 포함: O(서비스에 넣음) / X(뺌)',
    '- 나이: 제한 없음 / 아동·청소년(18세 이하) / 성인(19세 이상)',
    '- 운동재활: O(운동·재활 바우처로 시설 찾기에 나옴) / X(그 밖 장애인 사회서비스로 안내)',
    '- 장애정도: 무관 / 심한 장애만 / 심하지 않은 장애만 (복지로 서비스)',
]:
    guide.append([line])
guide.column_dimensions['A'].width = 110
guide['A1'].font = Font(bold=True, size=14)


def sheet(title, rows):
    ws = wb.create_sheet(title)
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    widths = [14, 10, 16, 38, 26, 9, 16, 12, 11, 14, 9, 16, 12, 11, 14, 24, 50, 60]
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = ws.dimensions
    head_fill, fix_fill, flag_fill = PatternFill('solid', fgColor='0F6B5C'), PatternFill('solid', fgColor='FFF3C4'), PatternFill('solid', fgColor='FDE2D2')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    n = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row[10:15]:
            c.fill = fix_fill
        if row[0].value:
            row[0].fill = flag_fill
        for c in row[16:18]:
            c.alignment = Alignment(wrap_text=False)
    for col, opts in (('K', 'O,X'), ('M', '제한 없음,아동·청소년,성인'), ('N', 'O,X'), ('O', '무관,심한 장애만,심하지 않은 장애만')):
        dv = DataValidation(type='list', formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{n}')
    dv = DataValidation(type='list', formula1='"전체,지체,뇌병변,시각,청각·언어,지적·자폐"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f'L2:L{n}')
    return ws


# 사회서비스: 지역사회서비스투자 서비스명별 1줄
P = json.load(open(os.path.join(RAW, 'ssis_providers.json'), encoding='utf-8'))
by = collections.defaultdict(list)
for r in P:
    if r.get('serviceTypeName') == '지역사회서비스투자':
        by[re.sub(r'\s+', ' ', r['serviceName'].strip())].append(r)
srows = []
for name, rs in sorted(by.items()):
    a = auto_social(name)
    rehab_word = bool(EXERCISE.search(name))
    if not a['include'] and not rehab_word:
        continue  # 장애와도 운동·재활과도 무관한 것은 표에서 뺌 (예: 산모 서비스)
    flag = ''
    if not a['include'] and rehab_word:
        flag = '빠졌는데 재활·운동 단어 있음'
    elif a['include'] and a['targets'] == '*' and not re.search(r'장애', name):
        flag = '대상 확인'
    regions = collections.Counter(re.sub(r'\s+', ' ', x.get('sidoName', '')) for x in rs)
    srows.append([flag, '사회서비스', name, name, f'{len(rs)}곳 · ' + ', '.join(f'{k} {v}' for k, v in regions.most_common(3)),
                  'O' if a['include'] else 'X', TG(a['targets']), AGE[a['age']], 'O' if a['ex'] else 'X', '', '', '', '', '', '', '', '', ''])
sheet('사회서비스', srows)

# 복지로: 서비스마다 1줄
W = json.load(open(os.path.join(RAW, 'welfare_disability.json'), encoding='utf-8'))
wrows = []
for level in ('central', 'local'):
    for r in W[level]:
        d = r.get('detail') or {}
        name = r['servNm'].strip()
        target = (d.get('tgtrDtlCn') or d.get('sprtTrgtCn') or '').strip()
        summary = (r.get('servDgst') or d.get('wlfareInfoOutlCn') or '').strip()
        a = auto_welfare(name, target, [], summary)
        # 장애인이 여러 대상 중 하나인 일반 서비스(에너지바우처 등)는 자동으로 넣어도 되므로 표시하지 않음.
        # 대상 원문에 중증 조건이 보이는데 자동으로 못 잡은 것만 표시
        flag = '장애정도 확인' if a['degree'] == '' and re.search(r'중증장애|장애의 정도가 심한|심한 장애', target) else ''
        where = '전국(중앙부처)' if level == 'central' else ' '.join(x for x in (r.get('ctpvNm', ''), r.get('sggNm', '')) if x)
        wrows.append([flag, '복지로', r['servId'], name, where, 'O', TG(a['targets']), AGE[a['age']], 'O' if a['ex'] else 'X', DEG[a['degree']],
                      '', '', '', '', '', '', summary[:300], target[:500]])
sheet('복지로', wrows)

wb.save(OUT)
print('사회서비스', len(srows), '줄 (먼저 볼 것', sum(1 for r in srows if r[0]), ') / 복지로', len(wrows), '줄 (먼저 볼 것', sum(1 for r in wrows if r[0]), ')')
