"""검수 끝난 엑셀 → pipeline/review/classification.csv (매주 자동 갱신이 읽는 파일)
사용: python read_review.py <검수표.xlsx>
"""
import csv, os, sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'review', 'classification.csv')
COLS = ['구분', '키', '서비스명', '확정_포함', '확정_대상장애유형', '확정_나이', '확정_운동재활', '확정_장애정도', '근거·메모']

wb = load_workbook(sys.argv[1], data_only=True)
out, n = [], 0
for title in ('사회서비스', '복지로'):
    ws = wb[title]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(head, row))
        if any((r.get(c) or '') for c in COLS[3:8]):
            out.append({c: (r.get(c) or '') for c in COLS})
            n += 1
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=COLS)
    w.writeheader()
    w.writerows(out)
print('검수 반영', n, '줄 →', OUT)
